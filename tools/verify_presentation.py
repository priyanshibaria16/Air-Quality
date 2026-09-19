"""
Read the generated deck back and check it against the brief's slide list.

    python tools/verify_presentation.py

Reports the title of every slide, how many bullets/shapes each holds, whether a
speaker note with a suggested timing exists, the total rehearsed length, and how
many pictures and tables were embedded. Fails if a slide is missing, if the order
differs from the brief's 17 titles, if a slide has no notes, or if the timings add
up to less than 10 or more than 15 minutes.
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE_TYPE  # noqa: E402

DECK = ROOT / "presentation" / "Air_Quality_Intelligence.pptx"

# The 17 slide titles the brief prescribes, in order (slide 1 is the cover).
BRIEF_SLIDES = ["Title", "Introduction", "Problem Statement", "Objectives", "Dataset",
                "Methodology", "Data Preprocessing", "EDA", "Correlation Analysis",
                "K-Means", "Anomaly Detection", "Classification",
                "Power BI Dashboard", "Results/Insights", "Limitations",
                "Future Scope", "Conclusion"]

# Titles as built, matched against the brief by keyword rather than string equality.
ACCEPTS = {
    "Title": ("Air Quality",), "Introduction": ("Introduction",),
    "Problem Statement": ("Problem Statement",), "Objectives": ("Objectives",),
    "Dataset": ("Dataset",), "Methodology": ("Methodology",),
    "Data Preprocessing": ("Preprocessing",), "EDA": ("Exploratory",),
    "Correlation Analysis": ("Correlation",), "K-Means": ("K-Means",),
    "Anomaly Detection": ("Anomaly",), "Classification": ("Classification",),
    "Power BI Dashboard": ("Power BI",), "Results/Insights": ("Results",),
    "Limitations": ("Limitations",), "Future Scope": ("Future Scope",),
    "Conclusion": ("Conclusion",),
}


def main() -> int:
    if not DECK.exists():
        print("MISSING", DECK, "- run tools/build_presentation.py first")
        return 1
    prs = Presentation(str(DECK))
    slides = list(prs.slides)
    print(f"slides: {len(slides)}   size: {prs.slide_width.inches:.2f} x "
          f"{prs.slide_height.inches:.2f} in   {DECK.stat().st_size:,} bytes")

    problems = []
    to_fill = []
    if len(slides) != len(BRIEF_SLIDES):
        problems.append(f"{len(slides)} slides, the brief asks for {len(BRIEF_SLIDES)}")

    total_minutes = 0.0
    for i, (slide, key) in enumerate(zip(slides, BRIEF_SLIDES), start=1):
        texts = []
        pictures = tables = 0
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pictures += 1
            if getattr(shape, "has_table", False) and shape.has_table:
                tables += 1
            if shape.has_text_frame:
                texts += [p.text.strip() for p in shape.text_frame.paragraphs
                          if p.text.strip()]
        title = texts[0] if texts else "(no text)"
        note = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
        timing = re.search(r"Suggested time: ([\d.]+) minutes", note)
        minutes = float(timing.group(1)) if timing else 0.0
        total_minutes += minutes
        ok = any(word in title for word in ACCEPTS.get(key, (key,)))
        print(f"  {i:>2}. {title[:52]:<52} pics={pictures} tables={tables} "
              f"lines={len(texts):>2} notes={len(note):>4}ch {minutes:>4.1f}min"
              + ("" if ok else "   <-- does not match the brief's "
                 f"{key!r} slide"))
        if not ok:
            problems.append(f"slide {i} should cover {key!r}")
        if len(slides) > 1 and not note:
            problems.append(f"slide {i} has no speaker notes")
        for text in texts:
            if re.search(r"<(STUDENT|ENROLMENT|INSTITUTION|\w+ NAME)", text):
                to_fill.append(f"slide {i}: {text[:60]}")

    if to_fill:
        print("front slide still to complete by hand:")
        for line in to_fill:
            print("   ", line)

    print(f"\nrehearsed length from the speaker notes: {total_minutes:.1f} minutes "
          "(brief: 10-15)")
    if not 10 <= total_minutes <= 15:
        problems.append(f"timings total {total_minutes:.1f} minutes, outside 10-15")

    print()
    if problems:
        print("PROBLEMS:")
        for problem in problems:
            print("  -", problem)
        return 1
    print("deck verified: the 17 brief slides are present in order, each carries "
          "speaker notes, and the timings total a 10-15 minute talk.")
    if to_fill:
        print("(the presenter's name/college on slide 1 are the only fields left to fill)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
